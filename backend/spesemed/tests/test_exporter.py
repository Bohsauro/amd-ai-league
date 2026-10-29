"""SPESEMED - Test offline per il modulo export Excel/ZIP (Fase 4B).

Copre:
  - `build_receipt_filename`: naming `YYYY-MM-DD_esercente_importoCAT.ext`
  - `export_excel`: genera .xlsx con fogli Riepilogo + Per Categoria
  - `export_zip` / `build_zip_bytes`: ZIP con Excel + ricevute rinominate
  - CLI `main`: flag `--no-excel` e guard contro ZIP vuoti

Nessuna dipendenza esterna (openpyxl è nel requirements; nessun network).
"""
from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest

from spesemed.core.exporter import (
    build_receipt_filename,
    build_zip_bytes,
    categoria_code,
    export_excel,
    export_zip,
    main,
)
from spesemed.core.models import SpesaMedica


# =============================================================================
# Fixtures / helpers
# =============================================================================

def _spesa(**overrides) -> SpesaMedica:
    base = dict(
        data_spesa="2024-03-15",
        esercente="Farmacia Centrale di Rossi Mario & C.",
        importo_totale=45.5,
        codice_fiscale_paziente="RSSMRA85A01H501W",
        categoria_spesa="Farmaci",
    )
    base.update(overrides)
    return SpesaMedica(**base)


@pytest.fixture
def spese() -> list[SpesaMedica]:
    return [
        _spesa(),
        _spesa(
            data_spesa="2024-04-02",
            esercente="Studio Dentistico Bianchi",
            importo_totale=120.0,
            categoria_spesa="Visita Specialistica",
        ),
        _spesa(
            data_spesa=None,
            esercente="Ottica Vision",
            importo_totale=80.0,
            categoria_spesa="Dispositivo Medico",
        ),
    ]


# =============================================================================
# Naming ricevute
# =============================================================================

class TestReceiptFilename:
    def test_pattern_completo(self):
        name = build_receipt_filename(_spesa())
        assert name == "2024-03-15_Farmacia-Centrale-di-Rossi-Mario-C_45.50EUR_FAR.pdf"

    def test_categoria_code(self):
        assert categoria_code("Farmaci") == "FAR"
        assert categoria_code("Visita Specialistica") == "VIS"
        assert categoria_code("Dispositivo Medico") == "DIS"
        assert categoria_code("Esami") == "ESA"
        assert categoria_code("Altro") == "ALT"
        assert categoria_code("Sconosciuta") == "ALT"  # fallback

    def test_estensione_preservata(self):
        name = build_receipt_filename(_spesa(), original_suffix=".jpg")
        assert name.endswith(".jpg")
        assert ".pdf" not in name

    def test_campi_mancanti_fallback(self):
        name = build_receipt_filename(
            _spesa(data_spesa=None, esercente=None, importo_totale=None)
        )
        assert name.startswith("senza-data_sconosciuto_0.00EUR_")

    def test_sanitizzazione_esercente(self):
        name = build_receipt_filename(_spesa(esercente="Farmacia//Rossi::& SRL"))
        # Solo caratteri [A-Za-z0-9 ._ -] nel nome esercente.
        assert "/" not in name
        assert ":" not in name


# =============================================================================
# Excel
# =============================================================================

class TestExportExcel:
    def test_genera_xlsx_con_fogli(self, spese, tmp_path: Path):
        out = export_excel(spese, tmp_path / "riepilogo.xlsx")
        assert out.exists()
        assert out.suffix == ".xlsx"

        from openpyxl import load_workbook
        wb = load_workbook(out)
        assert "Riepilogo" in wb.sheetnames
        assert "Per Categoria" in wb.sheetnames

        ws = wb["Riepilogo"]
        # 1 header + 3 spese + 1 totale = 5 righe.
        assert ws.max_row == 5
        assert ws.cell(row=1, column=1).value == "Data"
        assert ws.cell(row=2, column=4).value == "Farmaci"
        # Riga totale (riga 5): importo = 45.5 + 120 + 80 = 245.5.
        assert ws.cell(row=5, column=2).value == "TOTALE"
        assert ws.cell(row=5, column=3).value == pytest.approx(245.5)

        ws_cat = wb["Per Categoria"]
        # Categorie presenti: 3 (Farmaci, Visita, Dispositivo).
        cats = [ws_cat.cell(row=r, column=1).value for r in range(2, 5)]
        assert set(cats) == {"Farmaci", "Visita Specialistica", "Dispositivo Medico"}

    def test_xlsx_senza_spese(self, tmp_path: Path):
        out = export_excel([], tmp_path / "vuoto.xlsx")
        assert out.exists()
        from openpyxl import load_workbook
        wb = load_workbook(out)
        ws = wb["Riepilogo"]
        # Solo header + totale (0 spese).
        assert ws.max_row == 2
        assert ws.cell(row=2, column=3).value == 0.0


# =============================================================================
# ZIP
# =============================================================================

class TestExportZip:
    def test_zip_con_excel_e_ricevute(self, spese, tmp_path: Path):
        receipts = {
            "IMG_20240115.jpg": b"fake-jpg-content",
            "Scan2024-04-02.pdf": b"fake-pdf-content",
            "ricevuta-ottica.png": b"fake-png-content",
        }
        out = export_zip(spese, receipts, tmp_path / "export.zip")
        assert out.exists()

        with zipfile.ZipFile(out) as zf:
            names = zf.namelist()
            assert "riepilogo_spese.xlsx" in names
            # 3 ricevute rinominate + excel = 4 entry.
            assert len(names) == 4
            renamed = [n for n in names if n.startswith("ricevute/")]
            assert renamed == [
                "ricevute/2024-03-15_Farmacia-Centrale-di-Rossi-Mario-C_45.50EUR_FAR.jpg",
                "ricevute/2024-04-02_Studio-Dentistico-Bianchi_120.00EUR_VIS.pdf",
                "ricevute/senza-data_Ottica-Vision_80.00EUR_DIS.png",
            ]
            # Excel dentro lo ZIP è un xlsx valido e identico all'export
            # standalone (entrambi i fogli: Riepilogo + Per Categoria).
            from openpyxl import load_workbook
            excel_bytes = zf.read("riepilogo_spese.xlsx")
            assert excel_bytes[:2] == b"PK"
            wb_in_zip = load_workbook(io.BytesIO(excel_bytes))
            assert set(wb_in_zip.sheetnames) == {"Riepilogo", "Per Categoria"}

    def test_build_zip_bytes_roundtrip(self, spese):
        receipts = {"ric.pdf": b"x" * 100}
        data = build_zip_bytes(spese[:1], receipts)
        assert data[:2] == b"PK"  # magic ZIP

        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            assert zf.read("ricevute/2024-03-15_Farmacia-Centrale-di-Rossi-Mario-C_45.50EUR_FAR.pdf") == b"x" * 100

    def test_zip_spese_senza_ricevute(self, spese, tmp_path: Path):
        """Excel-only ZIP: spese senza ricevute corrispondenti."""
        out = export_zip(spese, {}, tmp_path / "solo-excel.zip")
        assert out.exists()
        with zipfile.ZipFile(out) as zf:
            assert zf.namelist() == ["riepilogo_spese.xlsx"]

    def test_zip_nomi_duplicati_dedup(self, tmp_path: Path):
        """Due ricevute identiche (stessa farmacia, stesso giorno, stesso importo)
        non devono produrre entry duplicate: la seconda riceve suffisso _2."""
        spesa = _spesa()  # 2024-03-15 Farmacia ... 45.50EUR FAR
        receipts = {"ric1.pdf": b"a", "ric2.pdf": b"b"}
        out = export_zip([spesa, spesa], receipts, tmp_path / "dedup.zip")

        with zipfile.ZipFile(out) as zf:
            renamed = sorted(n for n in zf.namelist() if n.startswith("ricevute/"))
            assert renamed == [
                "ricevute/2024-03-15_Farmacia-Centrale-di-Rossi-Mario-C_45.50EUR_FAR.pdf",
                "ricevute/2024-03-15_Farmacia-Centrale-di-Rossi-Mario-C_45.50EUR_FAR_2.pdf",
            ]
            # Ogni entry è univoca e i contenuti non vengono confusi.
            assert zf.read(renamed[0]) == b"a"
            assert zf.read(renamed[1]) == b"b"


# =============================================================================
# CLI
# =============================================================================

class TestCli:
    def _spese_json(self, tmp_path: Path) -> Path:
        p = tmp_path / "spese.json"
        p.write_text(
            json.dumps([_spesa().model_dump()]),
            encoding="utf-8",
        )
        return p

    def test_cli_xlsx(self, tmp_path: Path):
        """Output .xlsx → genera solo Excel."""
        rc = main([str(self._spese_json(tmp_path)), "-o", str(tmp_path / "out.xlsx")])
        assert rc == 0
        assert (tmp_path / "out.xlsx").exists()

    def test_cli_zip_con_receipts(self, tmp_path: Path):
        """ZIP con Excel + ricevute da --receipts-dir."""
        receipts_dir = tmp_path / "ric"
        receipts_dir.mkdir()
        (receipts_dir / "scan.pdf").write_bytes(b"pdf")
        rc = main([
            str(self._spese_json(tmp_path)),
            "-o", str(tmp_path / "out.zip"),
            "--receipts-dir", str(receipts_dir),
        ])
        assert rc == 0
        with zipfile.ZipFile(tmp_path / "out.zip") as zf:
            assert "riepilogo_spese.xlsx" in zf.namelist()
            assert any(n.startswith("ricevute/") for n in zf.namelist())

    def test_cli_no_excel_senza_receipts_errore(self, tmp_path: Path):
        """--no-excel senza --receipts-dir → errore (ZIP vuoto), exit 1."""
        rc = main([
            str(self._spese_json(tmp_path)),
            "-o", str(tmp_path / "out.zip"),
            "--no-excel",
        ])
        assert rc == 1
        assert not (tmp_path / "out.zip").exists()

    def test_cli_no_excel_con_receipts(self, tmp_path: Path):
        """--no-excel con --receipts-dir → ZIP solo ricevute (niente Excel)."""
        receipts_dir = tmp_path / "ric"
        receipts_dir.mkdir()
        (receipts_dir / "scan.pdf").write_bytes(b"pdf")
        rc = main([
            str(self._spese_json(tmp_path)),
            "-o", str(tmp_path / "out.zip"),
            "--receipts-dir", str(receipts_dir),
            "--no-excel",
        ])
        assert rc == 0
        with zipfile.ZipFile(tmp_path / "out.zip") as zf:
            assert "riepilogo_spese.xlsx" not in zf.namelist()
            assert any(n.startswith("ricevute/") for n in zf.namelist())

    def test_cli_no_excel_receipts_dir_vuoto_errore(self, tmp_path: Path):
        """--no-excel con --receipts-dir vuoto → errore (ZIP vuoto), exit 1."""
        empty_dir = tmp_path / "vuoto"
        empty_dir.mkdir()
        rc = main([
            str(self._spese_json(tmp_path)),
            "-o", str(tmp_path / "out.zip"),
            "--receipts-dir", str(empty_dir),
            "--no-excel",
        ])
        assert rc == 1
        assert not (tmp_path / "out.zip").exists()
